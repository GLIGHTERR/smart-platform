#!/usr/bin/env python3
"""Convert the four SmartChủ source documents to traceable Markdown mirrors.

The source documents are historical baselines. They intentionally retain old
phone/SMS authentication text and business rules. Current Markdown is written
to the canonical ``docs/`` tree. Baselines replaced by approved living specs
are written to ``archieve/source-derived`` for traceability.
"""

from __future__ import annotations

import hashlib
import re
import shutil
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATE = "2026-09-30"

SRS_SOURCE = ROOT / "archieve" / "legacy-source" / "SRS" / "SmartChu" / "SRS (SmartChủ).docx"
FRS_SOURCE = ROOT / "archieve" / "legacy-source" / "FRS" / "SmartChu" / "FRS - SmartChủ.docx"
REQ_SOURCE = ROOT / "archieve" / "legacy-source" / "Requirement_List" / "SmartChu" / "Requirements List - SmartChủ.xlsx"
US_SOURCE = ROOT / "archieve" / "legacy-source" / "User_Stories" / "SmartChu" / "User Story - SmartChủ.xlsx"

SRS_OUT = ROOT / "docs" / "SRS" / "SmartChu" / "SRS.md"
SRS_UC_DIR = ROOT / "docs" / "SRS" / "SmartChu" / "use-cases"
SRS_ARCHIVE_UC_DIR = ROOT / "archieve" / "source-derived" / "SRS" / "SmartChu" / "use-cases"
FRS_OUT = ROOT / "docs" / "FRS" / "SmartChu" / "FRS.md"
REQ_OUT = ROOT / "docs" / "Requirement_List" / "SmartChu" / "Requirements.md"
US_OUT = ROOT / "docs" / "User_Stories" / "SmartChu" / "User-Stories.md"
SRS_ASSET_DIR = ROOT / "docs" / "SRS" / "SmartChu" / "assets"
FRS_ASSET_DIR = ROOT / "docs" / "FRS" / "SmartChu" / "assets" / "source"

EXPECTED_HASHES = {
    SRS_SOURCE: "fc374c30bc5c6777cb4bd7f6d8e48988b9dc801bd6b203d72fbd19b4469b58ab",
    FRS_SOURCE: "c45b73e21d9ef7e090ec42b0620b5575792d36a93d21867f4b81947222af4181",
    REQ_SOURCE: "b329e93bacc322889718c2cbbc9f6f27833975eccaeb29ed50f246e415ef5e44",
    US_SOURCE: "6d061232942652055cb034b0908f7a702274e737574abb79da5f100098f7c374",
}

UC_AMENDMENTS = {
    1: "UC-01-sign-up.md",
    2: "UC-02-sign-in.md",
    3: "UC-03-forgot-password.md",
    6: "UC-06-view-business-information.md",
    7: "UC-07-submit-business-verification.md",
    10: "UC-10-create-electronic-contract.md",
    18: "UC-18-create-property.md",
    22: "UC-22-create-room.md",
}
POLICY_UCS = {6, 7, 10, 18, 22}
ARCHIVE_POLICY_LINK = "../../../../supplemental/Business_Rules/SmartChu/11-smartchu-business-verification-policy.md"
REQUIREMENT_TO_UC = {
    "SM038": 1,
    "SM039": 2,
    "SM041": 3,
    "SM045": 6,
    "SM046": 7,
    "SM048": 10,
    "SM055": 18,
    "SM059": 22,
}

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"


@dataclass(frozen=True)
class CellParagraph:
    text: str
    numbered: bool


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_sources() -> dict[Path, str]:
    actual: dict[Path, str] = {}
    for path, expected in EXPECTED_HASHES.items():
        value = sha256(path)
        if value != expected:
            raise RuntimeError(f"Source hash changed for {path}: {value}")
        actual[path] = value
    return actual


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    ascii_value = ascii_value.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def md_escape(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", "<br>").strip()


def normalize_id(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def xml_paragraph_text(paragraph: ET.Element) -> str:
    chunks: list[str] = []
    for node in paragraph.iter():
        if node.tag == W + "t":
            chunks.append(node.text or "")
        elif node.tag == W + "tab":
            chunks.append("\t")
        elif node.tag == W + "br":
            chunks.append("\n")
    return "".join(chunks).strip()


def xml_cell_paragraphs(cell: ET.Element) -> list[CellParagraph]:
    result: list[CellParagraph] = []
    for paragraph in cell.iter(W + "p"):
        text = xml_paragraph_text(paragraph)
        if not text:
            continue
        properties = paragraph.find(W + "pPr")
        numbered = properties is not None and properties.find(W + "numPr") is not None
        result.append(CellParagraph(text=text, numbered=numbered))
    return result


def xml_table_rows(table: ET.Element) -> list[list[list[CellParagraph]]]:
    return [
        [xml_cell_paragraphs(cell) for cell in row.findall(W + "tc")]
        for row in table.findall(W + "tr")
    ]


def plain_cell(paragraphs: list[CellParagraph]) -> str:
    return " / ".join(item.text for item in paragraphs)


def extract_srs_tables() -> tuple[list[tuple[str, str, str]], list[dict[str, list[CellParagraph]]]]:
    with ZipFile(SRS_SOURCE) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))

    catalogue: list[tuple[str, str, str]] = []
    details: list[dict[str, list[CellParagraph]]] = []
    for table in root.iter(W + "tbl"):
        rows = xml_table_rows(table)
        if not rows:
            continue
        header = [plain_cell(cell) for cell in rows[0]]
        if header[:3] == ["STT", "Mã UC", "Tên UC"]:
            for row in rows[1:]:
                if len(row) >= 3:
                    values = tuple(plain_cell(cell) for cell in row[:3])
                    if all(values):
                        catalogue.append(values)
        elif header[:2] == ["Nội dung", "Mô tả"]:
            fields: dict[str, list[CellParagraph]] = {}
            for row in rows[1:]:
                if len(row) >= 2:
                    key = plain_cell(row[0]).strip()
                    if key:
                        fields[key] = row[1]
            if fields:
                details.append(fields)

    if len(catalogue) != 46 or len(details) != 46:
        raise RuntimeError(
            f"Unexpected SmartChủ SRS counts: catalogue={len(catalogue)}, details={len(details)}"
        )
    return catalogue, details


def render_field(paragraphs: list[CellParagraph]) -> list[str]:
    if not paragraphs:
        return ["_Không có nội dung trong nguồn._"]
    if len(paragraphs) == 1:
        return [paragraphs[0].text]
    return [f"{'1.' if item.numbered else '-'} {item.text}" for item in paragraphs]


def uc_number(fields: dict[str, list[CellParagraph]]) -> int:
    raw = plain_cell(fields.get("Use Case ID", []))
    match = re.search(r"(\d+)", raw)
    if not match:
        raise RuntimeError(f"Cannot parse UC number from {raw!r}")
    return int(match.group(1))


def uc_filename(number: int, name: str) -> str:
    return f"UC-{number:02d}-{slugify(name)}.md"


def amendment_block(number: int) -> list[str]:
    lines: list[str] = []
    amendment = UC_AMENDMENTS.get(number)
    if amendment:
        path = f"../../../../../docs/SRS/SmartChu/use-cases/{amendment}"
        lines.extend(
            [
                "## Amendment / living implementation spec hiện hành",
                "",
                f"- [Living spec UC-{number:02d}]({path})",
                "- Nội dung nguồn bên dưới được giữ nguyên để truy vết. Khi triển khai và kiểm thử, quyết định Approved trong living spec được ưu tiên nếu có khác biệt.",
            ]
        )
        if number in POLICY_UCS:
            lines.append(f"- [Chính sách xác minh kinh doanh SmartChủ]({ARCHIVE_POLICY_LINK})")
        lines.append("")
    return lines


def render_uc(fields: dict[str, list[CellParagraph]], source_hash: str) -> tuple[str, str]:
    number = uc_number(fields)
    name = plain_cell(fields.get("Use Case Name", [])).strip() or f"UC-{number}"
    lines = [
        f"# UC-{number:02d}: {name}",
        "",
        "> **Loại tài liệu:** Baseline chuyển đổi nguyên trạng từ SRS SmartChủ lịch sử.",
        f"> **Nguồn:** `archieve/legacy-source/SRS/SmartChu/SRS (SmartChủ).docx` — SHA-256 `{source_hash}` — chuyển đổi {DATE}.",
        "> **Lưu ý:** Tài liệu này không tự động phản ánh quyết định PO mới nếu file DOCX nguồn chưa được sửa.",
        "",
        *amendment_block(number),
        "## Đặc tả nguồn",
        "",
    ]
    preferred = [
        "Use Case ID",
        "Use Case Name",
        "Created By",
        "Last Updated By",
        "Date Created",
        "Date Last Updated",
        "Actors",
        "Description",
        "Trigger",
        "Preconditions",
        "Postconditions",
        "Normal Flow",
        "Alternative Flows",
        "Exceptions",
        "Priority",
        "Frequency of Use",
        "Business Rules",
        "Other Information",
        "Assumptions",
    ]
    ordered = [key for key in preferred if key in fields]
    ordered.extend(key for key in fields if key not in ordered)
    for key in ordered:
        lines.extend([f"### {key}", "", *render_field(fields[key]), ""])
    return uc_filename(number, name), "\n".join(lines).rstrip() + "\n"


def extract_docx_media(source: Path, output_dir: Path) -> list[Path]:
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    extracted: list[Path] = []
    with ZipFile(source) as archive:
        names = sorted(name for name in archive.namelist() if name.startswith("word/media/"))
        for index, name in enumerate(names, 1):
            suffix = Path(name).suffix.lower()
            target = output_dir / f"image-{index:02d}{suffix}"
            target.write_bytes(archive.read(name))
            extracted.append(target)
    return extracted


def paragraph_markdown(paragraph: DocxParagraph) -> list[str]:
    text = paragraph.text.strip()
    if not text:
        return []
    style = paragraph.style.name if paragraph.style else ""
    match = re.match(r"Heading\s+(\d+)", style, flags=re.I)
    if match:
        level = min(int(match.group(1)) + 1, 6)
        return [f"{'#' * level} {text}", ""]
    if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
        return [f"- {text}"]
    return [text, ""]


def docx_table_markdown(table: Table) -> list[str]:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows or not any(any(cell for cell in row) for row in rows):
        return []
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    lines = [
        "| " + " | ".join(md_escape(cell) for cell in padded[0]) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for row in padded[1:]:
        lines.append("| " + " | ".join(md_escape(cell) for cell in row) + " |")
    return lines + [""]


def render_docx_body(source: Path, title: str, source_hash: str) -> list[str]:
    doc = Document(source)
    lines = [
        f"# {title}",
        "",
        "> **Loại tài liệu:** Markdown mirror của tài liệu nguồn; không thay thế file DOCX để kiểm tra định dạng gốc.",
        f"> **Nguồn:** `{source.relative_to(ROOT)}` — SHA-256 `{source_hash}` — chuyển đổi {DATE}.",
        "",
    ]
    for item in doc.iter_inner_content():
        if isinstance(item, DocxParagraph):
            lines.extend(paragraph_markdown(item))
        elif isinstance(item, Table):
            lines.extend(docx_table_markdown(item))
    return lines


def write_srs(source_hash: str) -> None:
    catalogue, details = extract_srs_tables()
    SRS_UC_DIR.mkdir(parents=True, exist_ok=True)
    SRS_ARCHIVE_UC_DIR.mkdir(parents=True, exist_ok=True)
    protected = set(UC_AMENDMENTS.values())
    for stale in SRS_UC_DIR.glob("UC-*.md"):
        if stale.name not in protected:
            stale.unlink()
    for stale in SRS_ARCHIVE_UC_DIR.glob("UC-*.md"):
        stale.unlink()

    detail_index: dict[int, tuple[str, str]] = {}
    for fields in details:
        number = uc_number(fields)
        name = plain_cell(fields.get("Use Case Name", [])).strip()
        filename, content = render_uc(fields, source_hash)
        target_dir = SRS_ARCHIVE_UC_DIR if number in UC_AMENDMENTS else SRS_UC_DIR
        (target_dir / filename).write_text(content, encoding="utf-8")
        detail_index[number] = (name, filename)

    readme = [
        "# SmartChủ SRS — Danh mục Use Case chuẩn",
        "",
        f"> Khôi phục {len(details)} bảng đặc tả từ `archieve/legacy-source/SRS/SmartChu/SRS (SmartChủ).docx` ngày {DATE}.",
        f"> SHA-256 nguồn: `{source_hash}`.",
        "> Đây là vị trí chuẩn duy nhất của Use Case SmartChủ; baseline bị living spec thay thế được lưu trong archive.",
        "",
        "| Mã UC | Tên trong danh mục nguồn | Tài liệu chuẩn | Trạng thái |",
        "| --- | --- | --- | --- |",
    ]
    for stt, raw_id, catalogue_name in catalogue:
        number_match = re.search(r"(\d+)", raw_id)
        if not number_match:
            raise RuntimeError(f"Invalid catalogue UC id: {raw_id}")
        number = int(number_match.group(1))
        detail_name, filename = detail_index[number]
        amendment = UC_AMENDMENTS.get(number)
        canonical_filename = amendment or filename
        status = "Living spec hiện hành" if amendment else "Baseline chuyển đổi hiện hành"
        readme.append(
            f"| UC-{number:02d} | {md_escape(catalogue_name)} | [{md_escape(detail_name)}]({canonical_filename}) | {status} |"
        )
    readme.extend(
        [
            "",
            "## Quy tắc sử dụng",
            "",
            "1. Mỗi UC chỉ có một file hiện hành trong thư mục này.",
            "2. Baseline bị living spec thay thế chỉ nằm trong `archieve/source-derived/SRS/SmartChu/use-cases/`.",
            "3. Không tạo cây `docs/use-cases/` hoặc `markdown/` song song.",
            "",
        ]
    )
    (SRS_UC_DIR / "README.md").write_text("\n".join(readme), encoding="utf-8")

    assets = extract_docx_media(SRS_SOURCE, SRS_ASSET_DIR)
    lines = render_docx_body(SRS_SOURCE, "SRS SmartChủ", source_hash)
    lines.extend(
        [
            "## Danh mục Use Case và đặc tả tách riêng",
            "",
            f"> DOCX chứa {len(catalogue)} UC và {len(details)} bảng đặc tả lồng. Toàn bộ đã được tách để agent có thể đọc theo từng UC.",
            "",
            "| STT | Mã UC | Tên UC | Bản đặc tả |",
            "| ---: | --- | --- | --- |",
        ]
    )
    for stt, raw_id, name in catalogue:
        number = int(re.search(r"(\d+)", raw_id).group(1))
        detail_name, filename = detail_index[number]
        filename = UC_AMENDMENTS.get(number, filename)
        lines.append(
            f"| {md_escape(stt)} | UC-{number:02d} | {md_escape(name)} | [{md_escape(detail_name)}](use-cases/{filename}) |"
        )
    lines.extend(["", "## Tài sản hình ảnh trích xuất", ""])
    for asset in assets:
        rel = asset.relative_to(SRS_OUT.parent).as_posix()
        lines.append(f"- [{asset.name}]({rel})")
    lines.append("")
    SRS_OUT.write_text("\n".join(lines), encoding="utf-8")


def xml_all_tables(source: Path) -> list[list[list[str]]]:
    with ZipFile(source) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    result: list[list[list[str]]] = []
    for table in root.iter(W + "tbl"):
        rows: list[list[str]] = []
        for row in table.findall(W + "tr"):
            rows.append([plain_cell(xml_cell_paragraphs(cell)) for cell in row.findall(W + "tc")])
        if rows and any(any(cell for cell in row) for row in rows):
            result.append(rows)
    return result


def write_frs(source_hash: str) -> None:
    assets = extract_docx_media(FRS_SOURCE, FRS_ASSET_DIR)
    lines = render_docx_body(FRS_SOURCE, "FRS SmartChủ", source_hash)
    lines.extend(
        [
            "## Phụ lục bảng nguồn",
            "",
            "> Phụ lục liệt kê toàn bộ bảng trong XML (bao gồm bảng lồng) để tránh bỏ sót nội dung khi converter DOCX chỉ thấy bảng cấp cao nhất.",
            "",
        ]
    )
    for index, rows in enumerate(xml_all_tables(FRS_SOURCE), 1):
        lines.extend([f"### Bảng nguồn {index}", ""])
        width = max(len(row) for row in rows)
        padded = [row + [""] * (width - len(row)) for row in rows]
        lines.extend(
            [
                "| " + " | ".join(md_escape(cell) for cell in padded[0]) + " |",
                "| " + " | ".join("---" for _ in range(width)) + " |",
            ]
        )
        for row in padded[1:]:
            lines.append("| " + " | ".join(md_escape(cell) for cell in row) + " |")
        lines.append("")
    lines.extend(["## Tài sản hình ảnh trích xuất", ""])
    for asset in assets:
        rel = asset.relative_to(FRS_OUT.parent).as_posix()
        lines.append(f"- [{asset.name}]({rel})")
    lines.append("")
    FRS_OUT.write_text("\n".join(lines), encoding="utf-8")


def actual_bounds(ws) -> tuple[int, int]:
    last_row = 0
    last_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                last_row = max(last_row, cell.row)
                last_col = max(last_col, cell.column)
    return last_row, last_col


def amendment_for_uc(number: int) -> list[str]:
    filename = UC_AMENDMENTS.get(number)
    if not filename:
        return []
    lines = [
        f"- Living spec hiện hành: [UC-{number:02d}](../../../docs/SRS/SmartChu/use-cases/{filename})"
    ]
    if number in POLICY_UCS:
        lines.append(
            "- Chính sách liên quan: [Xác minh kinh doanh SmartChủ](../../../archieve/supplemental/Business_Rules/SmartChu/11-smartchu-business-verification-policy.md)"
        )
    return lines


def write_requirements(source_hash: str) -> None:
    wb = load_workbook(REQ_SOURCE, data_only=False)
    lines = [
        "# Requirements List — SmartChủ",
        "",
        "> **Loại tài liệu:** Markdown mirror của workbook nguồn.",
        f"> **Nguồn:** `archieve/legacy-source/Requirement_List/SmartChu/Requirements List - SmartChủ.xlsx` — SHA-256 `{source_hash}` — chuyển đổi {DATE}.",
        "> Các requirement lịch sử về số điện thoại/SMS được giữ nguyên. Living spec Approved trong `docs/` được ưu tiên khi triển khai.",
        "",
    ]
    for ws in wb.worksheets:
        max_row, max_col = actual_bounds(ws)
        lines.extend([f"## Sheet: {ws.title}", ""])
        if max_row <= 1:
            lines.extend(["_Sheet nguồn chỉ có hàng tiêu đề, chưa có dữ liệu._", ""])
            continue
        headers = [str(ws.cell(1, col).value or f"Column {col}").strip() for col in range(1, max_col + 1)]
        function_index = next(
            (index for index, header in enumerate(headers) if "Requirements/ Functions" in header),
            2,
        )
        for row in range(2, max_row + 1):
            values = [ws.cell(row, col).value for col in range(1, max_col + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            identifier = normalize_id(values[0]) if values else f"Row {row}"
            title_value = values[function_index] if function_index < len(values) else None
            title = str(title_value).strip() if title_value not in (None, "") else "Requirement"
            lines.extend([f"### {identifier} — {title}", ""])
            amendment_lines = amendment_for_uc(REQUIREMENT_TO_UC[identifier]) if identifier in REQUIREMENT_TO_UC else []
            if amendment_lines:
                lines.extend(["> **Amendment hiện hành**", *amendment_lines, ""])
            lines.extend(["| Trường | Giá trị |", "| --- | --- |"])
            for header, value in zip(headers, values):
                if value not in (None, ""):
                    lines.append(f"| {md_escape(header)} | {md_escape(value)} |")
            lines.append("")
    REQ_OUT.write_text("\n".join(lines), encoding="utf-8")


def write_user_stories(source_hash: str) -> None:
    wb = load_workbook(US_SOURCE, data_only=False)
    lines = [
        "# User Story — SmartChủ",
        "",
        "> **Loại tài liệu:** Markdown mirror của workbook nguồn.",
        f"> **Nguồn:** `archieve/legacy-source/User_Stories/SmartChu/User Story - SmartChủ.xlsx` — SHA-256 `{source_hash}` — chuyển đổi {DATE}.",
        "> User story lịch sử về số điện thoại/SMS được giữ nguyên. Living spec Approved trong `docs/` được ưu tiên khi triển khai.",
        "",
    ]
    story_count = 0
    for ws in wb.worksheets:
        max_row, _ = actual_bounds(ws)
        lines.extend([f"## {ws.title}", ""])
        current: dict[str, object] | None = None
        stories: list[dict[str, object]] = []
        for row in range(2, max_row + 1):
            story_id = ws.cell(row, 1).value
            if story_id not in (None, ""):
                current = {
                    "id": normalize_id(story_id),
                    "name": ws.cell(row, 2).value or "",
                    "story": ws.cell(row, 3).value or "",
                    "acs": [],
                    "priority": ws.cell(row, 5).value or "",
                }
                stories.append(current)
                story_count += 1
            if current is not None:
                ac = ws.cell(row, 4).value
                if ac not in (None, ""):
                    current["acs"].append(str(ac).strip())
        for story in stories:
            identifier = str(story["id"])
            lines.extend(
                [
                    f"### US-{int(float(identifier)):02d}: {story['name']}",
                    "",
                    f"- **Priority:** {story['priority'] or 'Không ghi trong nguồn'}",
                    f"- **User Story:** {story['story']}",
                ]
            )
            amendment_lines = amendment_for_uc(int(float(identifier)))
            if amendment_lines:
                lines.extend(amendment_lines)
            lines.extend(["", "**Acceptance Criteria nguồn**", ""])
            for ac in story["acs"]:
                lines.append(f"- {ac}")
            lines.append("")
    if story_count != 46:
        raise RuntimeError(f"Unexpected SmartChủ user-story count: {story_count}")
    US_OUT.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    hashes = verify_sources()
    for path in [SRS_OUT.parent, FRS_OUT.parent, REQ_OUT.parent, US_OUT.parent]:
        path.mkdir(parents=True, exist_ok=True)
    write_srs(hashes[SRS_SOURCE])
    write_frs(hashes[FRS_SOURCE])
    write_requirements(hashes[REQ_SOURCE])
    write_user_stories(hashes[US_SOURCE])
    print("SmartChủ Markdown mirrors generated")
    print("srs_use_cases=46")
    print("user_stories=46")


if __name__ == "__main__":
    main()
